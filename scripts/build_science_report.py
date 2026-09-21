"""Validate saved FL experiments, draw figures, and write a Traditional Chinese Word report. No training."""

import argparse
import csv
from datetime import date
import hashlib
import json
import math
import os
from pathlib import Path
import statistics
import zipfile

os.environ.setdefault("MPLCONFIGDIR", "/tmp/fl-report-matplotlib")
import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


# 讀取既有實驗 CSV，這支報告腳本不負責訓練。
def readcsv(path):
    with path.open(newline="") as f:
        return list(csv.DictReader(f))


# 報告流程：核對資料 → 重算統計 → 畫圖 → 編排 Word → 檢查輸出。
# TISF.docx 只作章節參考，原始文件與實驗結果不會被改寫。
def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--study", type=Path, default=ROOT / "results/studies/alpha_scale"
    )
    parser.add_argument("--template", type=Path, default=ROOT / "TISF.docx")
    parser.add_argument("--output", type=Path, default=ROOT / "reports/current")
    args = parser.parse_args()
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    figures = out / "figures"
    figures.mkdir(exist_ok=True)
    # 只讀取參考文件的章節，不複製原題目的成果或作者資訊。
    template = Document(args.template)
    template_headings = [
        p.text
        for p in template.paragraphs
        if p.text.startswith(("壹、", "貳、", "參、", "肆、", "伍、", "陸、"))
    ]
    fontfile = Path("/mnt/c/Windows/Fonts/msjh.ttc")
    if fontfile.exists():
        font_manager.fontManager.addfont(str(fontfile))
        fontname = font_manager.FontProperties(fname=str(fontfile)).get_name()
    else:
        available = {f.name for f in font_manager.fontManager.ttflist}
        fontname = next(
            (
                f
                for f in ["Noto Sans CJK TC", "Microsoft JhengHei", "Noto Sans CJK JP"]
                if f in available
            ),
            None,
        )
        if not fontname:
            raise RuntimeError(
                "Install a Traditional Chinese font before generating figures."
            )
    plt.rcParams.update(
        {
            "font.family": fontname,
            "axes.unicode_minus": False,
            "font.size": 11,
            "axes.spines.top": False,
            "axes.spines.right": False,
        }
    )
    # 第一階段：依 study.json 找到原始結果，交會設定只計算一次。
    manifest = json.loads((args.study / "study.json").read_text())
    summary = readcsv(args.study / "summary.csv")
    per_seed = readcsv(args.study / "per_seed.csv")
    datasets, unique, weights, computed = {}, set(), [], []
    # 每列對應一個實驗設定；下方依 seed 逐一驗證。
    for entry in manifest["entries"]:
        folder = Path(entry["path"])
        if not folder.exists():
            folder = ROOT / str(folder).split("FL-project/", 1)[-1]
        if folder not in unique:
            for name, digest in json.loads(
                (folder / "checksums.json").read_text()
            ).items():
                assert (
                    hashlib.sha256((folder / name).read_bytes()).hexdigest() == digest
                ), (folder / name)
            unique.add(folder)
        records = []
        for seed in [42, 43, 44]:
            pair = {}
            cfgs = []
            for mode in ["clean", "poisoned"]:
                run = folder / f"seed_{seed}" / mode
                cfg = json.loads((run / "config.json").read_text())
                assert (
                    cfg["alpha"] == entry["alpha"]
                    and cfg["poison_scale"] == entry["scale"]
                )
                assert cfg["seed"] == seed and cfg["enable_poison"] == (
                    mode == "poisoned"
                )
                cfgs.append(cfg)
                rows = readcsv(run / "results.csv")
                assert [int(r["round"]) for r in rows] == list(range(1, 11))
                for r in rows:
                    assert (
                        math.isfinite(float(r["loss"]))
                        and 0 <= float(r["accuracy"]) <= 1
                    )
                pair[mode] = rows
                with np.load(run / "partitions.npz") as partitions:
                    parts = [partitions[f"client_{i}"] for i in range(5)]
                    assert np.array_equal(
                        np.sort(np.concatenate(parts)), np.arange(60000)
                    )
                    assert [len(p) for p in parts] == cfg["client_sample_counts"]
            for name in [
                "initial_parameters_sha256",
                "partition_sha256",
                "shuffle_seeds",
                "client_sample_counts",
            ]:
                assert cfgs[0][name] == cfgs[1][name]
            ca, pa = [float(pair[m][-1]["accuracy"]) for m in ["clean", "poisoned"]]
            cl, pl = [float(pair[m][-1]["loss"]) for m in ["clean", "poisoned"]]
            record = dict(
                experiment=entry["experiment"],
                alpha=entry["alpha"],
                scale=entry["scale"],
                seed=seed,
                clean_accuracy=ca,
                poisoned_accuracy=pa,
                clean_loss=cl,
                poisoned_loss=pl,
                accuracy_damage_pp=100 * (ca - pa),
                loss_damage=pl - cl,
            )
            matches = [
                r
                for r in per_seed
                if r["experiment"] == entry["experiment"]
                and float(r["alpha"]) == entry["alpha"]
                and float(r["scale"]) == entry["scale"]
                and int(r["seed"]) == seed
            ]
            assert len(matches) == 1
            for metric in [
                "clean_accuracy",
                "poisoned_accuracy",
                "clean_loss",
                "poisoned_loss",
                "accuracy_damage_pp",
                "loss_damage",
            ]:
                assert math.isclose(
                    float(matches[0][metric]), record[metric], abs_tol=1e-10
                )
            weight = cfgs[0]["client_sample_counts"][0] / 60000
            weights.append(
                {
                    **record,
                    "attacker_samples": cfgs[0]["client_sample_counts"][0],
                    "attacker_weight": weight,
                }
            )
            records.append(record)
            datasets[(entry["experiment"], entry["alpha"], entry["scale"], seed)] = pair
        summary_row = next(
            r
            for r in summary
            if r["experiment"] == entry["experiment"]
            and float(r["alpha"]) == entry["alpha"]
            and float(r["scale"]) == entry["scale"]
        )
        for metric in [
            "clean_accuracy",
            "poisoned_accuracy",
            "clean_loss",
            "poisoned_loss",
            "accuracy_damage_pp",
            "loss_damage",
        ]:
            for suffix, fn in [
                ("mean", statistics.mean),
                ("sample_std", statistics.stdev),
            ]:
                assert math.isclose(
                    float(summary_row[metric + "_" + suffix]),
                    fn(r[metric] for r in records),
                    abs_tol=1e-10,
                )
        computed.extend(records)
    with (out / "attacker_weight_and_damage.csv").open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(weights[0]))
        writer.writeheader()
        writer.writerows(weights)
    audit = {
        "template": str(args.template),
        "template_sections": template_headings,
        "unique_settings": len(unique),
        "training_runs": len(unique) * 6,
        "summary_rows": len(summary),
        "paired_rows_including_shared_setting": len(computed),
        "checks": "Raw checksums, full partitions, paired identities and summary statistics verified",
        "input_sha256": {
            str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in [
                args.template,
                args.study / "summary.csv",
                args.study / "per_seed.csv",
                args.study / "study.json",
            ]
        },
    }
    (out / "validation.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n"
    )

    # 第二階段：全部圖表由核對過的資料重建，不手動輸入曲線數值。
    colors = ["#1676a3", "#df762b", "#548c45"]

    # 同一張圖同時輸出 PNG 與 PDF，再關閉畫布以釋放記憶體。
    def save(fig, name):
        fig.savefig(figures / f"{name}.png", dpi=200, bbox_inches="tight")
        fig.savefig(figures / f"{name}.pdf", bbox_inches="tight")
        plt.close(fig)

    for ex, variable in [("A", "alpha"), ("B", "scale")]:
        rows = [r for r in summary if r["experiment"] == ex]
        x = np.arange(len(rows))
        labels = [f"{float(r[variable]):g}" for r in rows]
        fig, axes = plt.subplots(2, 2, figsize=(10, 7.3), layout="constrained")
        for ax, metric, ylabel, factor in [
            (axes[0, 0], "accuracy", "準確率（%）", 100),
            (axes[0, 1], "loss", "Loss（越低越好）", 1),
        ]:
            for mode, label, color in [
                ("clean", "Clean：無攻擊", colors[0]),
                ("poisoned", "Poisoned：開啟放大", colors[1]),
            ]:
                ax.errorbar(
                    x,
                    [float(r[f"{mode}_{metric}_mean"]) * factor for r in rows],
                    yerr=[
                        float(r[f"{mode}_{metric}_sample_std"]) * factor for r in rows
                    ],
                    fmt="o-",
                    capsize=4,
                    label=label,
                    color=color,
                )
            ax.set_ylabel(ylabel)
            ax.legend(fontsize=9)
        for ax, metric, label in [
            (axes[1, 0], "accuracy_damage_pp", "準確率損害（百分點）"),
            (axes[1, 1], "loss_damage", "Loss 增加量"),
        ]:
            ax.errorbar(
                x,
                [float(r[metric + "_mean"]) for r in rows],
                yerr=[float(r[metric + "_sample_std"]) for r in rows],
                fmt="o-",
                capsize=4,
                color="#73499d",
            )
            ax.axhline(0, color="gray", lw=1, ls="--")
            ax.set_ylabel(label)
        for ax in axes.flat:
            ax.set_xticks(x, labels)
            ax.set_xlabel(
                "Alpha（依設定等距排列）" if ex == "A" else "Scale（依設定等距排列）"
            )
            ax.grid(alpha=0.2)
        fig.suptitle(
            f"實驗 {ex}｜"
            + ("固定 Scale = 10" if ex == "A" else "固定 Alpha = 0.1")
            + "\n點為三個 seeds 的平均；誤差棒為樣本標準差（非信賴區間）",
            fontsize=13,
        )
        save(fig, f"{ex}_overview")
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.2), layout="constrained")
    for ax, ex, var in zip(axes, ["A", "B"], ["alpha", "scale"]):
        rows = [r for r in summary if r["experiment"] == ex]
        x = np.arange(len(rows))
        for seed, color in zip([42, 43, 44], colors):
            data = [r for r in computed if r["experiment"] == ex and r["seed"] == seed]
            ax.plot(
                x,
                [r["accuracy_damage_pp"] for r in data],
                "o-",
                label=f"Seed {seed}",
                color=color,
            )
        ax.set_xticks(x, [f"{float(r[var]):g}" for r in rows])
        ax.set_xlabel(var.title() + "（依設定等距排列）")
        ax.set_ylabel("準確率損害（百分點）")
        ax.set_title(f"實驗 {ex}：各 seed 的結果")
        ax.axhline(0, color="gray", ls="--")
        ax.grid(alpha=0.2)
        ax.legend()
    save(fig, "seed_damage")
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.2), layout="constrained")
    for seed, color in zip([42, 43, 44], colors):
        data = [r for r in weights if r["experiment"] == "A" and r["seed"] == seed]
        axes[0].plot(
            range(5),
            [r["attacker_weight"] * 100 for r in data],
            "o-",
            color=color,
            label=f"Seed {seed}",
        )
        axes[1].scatter(
            [r["attacker_weight"] * 100 for r in data],
            [r["accuracy_damage_pp"] for r in data],
            color=color,
            label=f"Seed {seed}",
        )
    axes[0].set_xticks(range(5), ["0.01", "0.1", "1", "10", "100"])
    axes[0].set_xlabel("Alpha（依設定等距排列）")
    axes[0].set_ylabel("攻擊者樣本占比（%）")
    axes[1].set_xlabel("攻擊者樣本占比（%）")
    axes[1].set_ylabel("準確率損害（百分點）")
    for ax in axes:
        ax.grid(alpha=0.2)
        ax.legend()
    fig.suptitle(
        "實驗 A：攻擊者權重與損害｜散點只能顯示關係，不能證明因果", fontsize=13
    )
    save(fig, "attacker_weight")
    fig, axes = plt.subplots(1, 2, figsize=(10, 4.2), layout="constrained")
    for scale in [1, 2, 5, 10, 20]:
        values = [
            datasets[("B", 0.1, scale, seed)]["poisoned"] for seed in [42, 43, 44]
        ]
        for ax, metric, factor in [(axes[0], "accuracy", 100), (axes[1], "loss", 1)]:
            y = np.array([[float(r[metric]) * factor for r in rows] for rows in values])
            ax.plot(range(1, 11), y.mean(axis=0), marker=".", label=f"Scale {scale}")
    for ax, metric, factor in [(axes[0], "accuracy", 100), (axes[1], "loss", 1)]:
        y = np.array(
            [
                [float(r[metric]) * factor for r in datasets[("B", 0.1, 1, s)]["clean"]]
                for s in [42, 43, 44]
            ]
        )
        ax.plot(range(1, 11), y.mean(axis=0), "k--", label="Clean")
        ax.set_xlabel("聯邦訓練輪數")
        ax.grid(alpha=0.2)
        ax.legend(fontsize=8)
    axes[0].set_ylabel("平均準確率（%）")
    axes[1].set_ylabel("平均 Loss")
    fig.suptitle("實驗 B：十輪訓練變化（僅畫平均；變異見總覽與各 seed 圖）")
    save(fig, "training_curves")
    fig, ax = plt.subplots(figsize=(8, 6))
    ax.axis("off")
    boxes = [
        (0.5, 0.94, "研究：資料分布與更新放大對聯邦學習的影響"),
        (0.24, 0.74, "實驗 A：改 Alpha\n0.01 / 0.1 / 1 / 10 / 100\n固定 Scale = 10"),
        (0.76, 0.74, "實驗 B：改 Scale\n1 / 2 / 5 / 10 / 20\n固定 Alpha = 0.1"),
        (0.5, 0.48, "每個設定使用 seeds 42 / 43 / 44"),
        (0.25, 0.28, "Clean\n正常訓練"),
        (0.75, 0.28, "Poisoned\nClient 0 放大更新"),
        (0.5, 0.07, "比較 Accuracy / Loss → 計算配對攻擊損害"),
    ]
    for x, y, t in boxes:
        ax.text(
            x,
            y,
            t,
            ha="center",
            va="center",
            fontsize=11,
            bbox=dict(boxstyle="round,pad=.55", fc="#edf4fa", ec="#35759a"),
        )
    for start, end in [
        ((0.43, 0.87), (0.25, 0.83)),
        ((0.57, 0.87), (0.75, 0.83)),
        ((0.24, 0.63), (0.44, 0.53)),
        ((0.76, 0.63), (0.56, 0.53)),
        ((0.44, 0.42), (0.25, 0.35)),
        ((0.56, 0.42), (0.75, 0.35)),
        ((0.25, 0.20), (0.43, 0.12)),
        ((0.75, 0.20), (0.57, 0.12)),
    ]:
        ax.annotate(
            "",
            xy=end,
            xytext=start,
            arrowprops=dict(arrowstyle="->", color="#426679", lw=1.5),
        )
    save(fig, "study_flow")

    # 第三階段：設定 A4、字型與段落樣式，再逐頁加入結果。
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    for sty in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
        style = doc.styles[sty]
        style.font.name = "Microsoft JhengHei"
        style._element.get_or_add_rPr().rFonts.set(
            qn("w:eastAsia"), "Microsoft JhengHei"
        )
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Heading 1", 17), ("Heading 2", 13)]:
        doc.styles[name].font.size = Pt(size)
        doc.styles[name].font.color.rgb = RGBColor.from_string("205A78")
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("聯邦學習研究｜階段報告　•　")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    # 新增一般內文段落。
    def p(t):
        return doc.add_paragraph(t)

    # 新增主要章節標題。
    def h(t):
        doc.add_heading(t, 1)

    # 新增次要章節標題。
    def sub(t):
        doc.add_heading(t, 2)

    # 換到下一頁並加入章節標題，讓報告各段落更容易閱讀。
    def page(t):
        doc.add_page_break()
        h(t)

    # 建立表格，讓欄名跨頁重複，並避免同一列被切成兩頁。
    def table(headers, rows):
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = "Light Shading Accent 1"
        for c, t in zip(tb.rows[0].cells, headers):
            c.text = str(t)
        repeat = OxmlElement("w:tblHeader")
        tb.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = tb.add_row().cells
            for c, t in zip(cells, row):
                c.text = str(t)
        for row in tb.rows:
            prop = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(prop)
            for cell in row.cells:
                for para in cell.paragraphs:
                    para.paragraph_format.space_after = Pt(3)
                    for run in para.runs:
                        run.font.size = Pt(9)
        return tb

    # 插入已產生的圖檔與圖說；圖表來源仍是原始 CSV。
    def picture(name, caption, width=6.5):
        doc.add_picture(str(figures / f"{name}.png"), width=Inches(width))
        doc.add_paragraph(caption, "Caption")

    # 加入 Word 公式節點，並在下一段用文字解釋符號。
    def formula(text, explanation):
        para = doc.add_paragraph()
        para.alignment = 1
        mathnode = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        txt = OxmlElement("m:t")
        txt.text = text
        run.append(txt)
        mathnode.append(run)
        para._p.append(mathnode)
        p(explanation)

    p("科學展覽研究報告｜目前成果")
    doc.add_paragraph("資料分布與模型更新放大\n對聯邦學習的影響", style="Title")
    doc.add_paragraph("以 MNIST、FedAvg 與三個隨機種子進行比較", style="Subtitle")
    p("科別：電腦科學與資訊工程科")
    p("關鍵詞：聯邦學習、資料分布不均、模型更新放大、攻擊損害")
    p(f"整理日期：{date.today().isoformat()}")
    p("學校／作者／指導老師／參賽編號：待填")
    p(
        "本報告參考 TISF.docx 的章節安排，以目前完成的實驗資料撰寫。屬於階段成果，尚未完成防禦方法，也尚未進行後門或 Sybil 攻擊實驗。"
    )
    picture(
        "study_flow",
        "研究架構：兩組實驗分別改變一個設定，再以三個 seeds 比較 Clean 與 Poisoned。",
        5.5,
    )
    page("摘要")
    p(
        "本研究想了解：當五個參與端持有不同的資料時，聯邦學習受到模型更新放大攻擊後，辨識能力會如何改變。我們使用 MNIST 手寫數字資料，讓每個參與端先自行訓練，再以 FedAvg 合併模型。攻擊者不修改圖片或標籤，而是把自己的模型更新乘上指定倍率。"
    )
    p(
        "實驗 A 固定放大倍率 Scale = 10，改變 Alpha 為 0.01、0.1、1、10、100；實驗 B 固定 Alpha = 0.1，改變 Scale 為 1、2、5、10、20。每個設定使用 seeds 42、43、44，分別執行有攻擊與無攻擊的十輪訓練。兩組共九個不同設定、54 次訓練，交會設定只使用同一批結果。"
    )
    p(
        "結果顯示，實驗 A 的無攻擊平均準確率由 69.96% 上升至 90.93%；但平均攻擊損害也由 26.76 增加至 64.68 個百分點。實驗 B 的 Scale = 1 控制組準確率差為零；Scale 增加至 20 時，平均準確率損害為 40.05 個百分點。部分設定在三個 seeds 間有很大的差異，因此不能只用平均值判斷。Alpha 同時影響資料種類及攻擊者樣本數，後續需固定攻擊者權重，才能進一步分辨原因。"
    )
    sub("Abstract")
    p(
        "This study examines how client data distribution and update scaling affect federated learning on MNIST. Five clients train with FedAvg for ten rounds. Experiment A varies alpha with scale fixed at 10; Experiment B varies scale with alpha fixed at 0.1. Each setting uses three matched seeds with clean and scaled-update conditions. Nine unique settings produce 54 training runs. Mean accuracy damage increases from 26.76 to 64.68 percentage points across the alpha sweep and from 0 to 40.05 points across the scale sweep. Large differences between seeds remain. Since alpha also changes the attacker’s sample count and aggregation weight, these observations do not isolate the effect of label imbalance."
    )
    page("壹、前言")
    sub("一、研究動機")
    p(
        "聯邦學習的想法是讓資料留在各參與端，由各端訓練後傳回模型，再合併成共同模型［1］。例如，不同使用者可能擁有不同類型的資料。即使模型架構相同，學到的更新也可能不同。本研究關心的是：如果其中一個參與端把更新刻意放大，是否會影響整體模型？這種影響又是否與資料分布有關？"
    )
    sub("二、研究目的")
    for t in [
        "比較 Alpha 改變時，無攻擊與有攻擊模型的準確率及 Loss。",
        "比較 Scale 改變時，模型更新放大造成多少損害。",
        "以相同 seed 配對比較，再觀察三個 seeds 的差異。",
        "記錄攻擊者資料量，找出值得進一步控制與驗證的因素。",
    ]:
        p("• " + t)
    sub("三、名詞說明")
    table(
        ["名詞", "本研究中的意思"],
        [
            ["Client（參與端）", "持有部分訓練資料並訓練模型的角色。"],
            ["Non-IID", "各端資料分布不同，例如某端主要是數字 2，另一端主要是數字 8。"],
            [
                "Alpha（α）",
                "資料分配參數；小值通常使各端數字種類差異更大，大值通常較接近。",
            ],
            ["Seed", "決定隨機起始設定的編號；較大的編號不代表較強的攻擊。"],
            ["Scale（s）", "攻擊者模型更新的放大倍率。"],
            [
                "Clean／Poisoned",
                "無攻擊／開啟更新放大的條件；Scale = 1 是正常更新控制。",
            ],
            ["Loss", "衡量預測錯誤程度的分數；本研究使用交叉熵，通常越低越好。"],
        ],
    )
    page("貳、研究設備及器材")
    env = json.loads((next(iter(unique)) / "environment.json").read_text())
    table(
        ["項目", "設定或用途"],
        [
            [
                "資料集",
                "MNIST；60,000 張訓練圖片、10,000 張測試圖片；28 × 28 灰階、10 類數字。",
            ],
            ["模型", "攤平為 784 個輸入 → 128 個隱藏節點 → ReLU → 10 類輸出。"],
            [
                "執行方式",
                "CPU 單執行緒，依序模擬五個 client；使用 Flower 的加權合併函式。",
            ],
            ["Python", env["python"]],
            *[[k, v] for k, v in env["packages"].items()],
            ["繪圖與文件", "Matplotlib 畫圖，python-docx 產生 Word 報告。"],
        ],
    )
    p(
        "本報告使用各實驗保存的環境資料。未記錄於執行檔案中的 CPU 型號、記憶體容量及執行時間不自行填入；若正式送件需要硬體規格，應由作者補上。"
    )
    p(
        "測試資料不參與模型訓練。每輪只對共同測試集評估一次，因為五個 client 使用相同的全域模型和測試資料，無須重複算五次。這是本機模擬，不能用來判斷真實網路速度或通訊成本。MNIST 的載入方式參考 torchvision 文件［2］。"
    )
    sub("資料來源與使用範圍")
    p(
        "分析使用 results/studies/alpha_scale/ 的 study.json、per_seed.csv、summary.csv，並依 study.json 找到每次訓練的原始 results.csv 及 config.json。舊版含有參數快照錯誤的 poisoning 結果不納入本報告；較早的單次 baseline 也不混入三個 seeds 的統計。"
    )
    page("參、研究過程及方法")
    sub("一、實驗設定")
    table(
        ["實驗", "改變的設定", "固定的設定"],
        [
            ["A", "Alpha：0.01、0.1、1、10、100", "Scale = 10"],
            ["B", "Scale：1、2、5、10、20", "Alpha = 0.1"],
        ],
    )
    table(
        ["共同條件", "設定"],
        [
            ["Seeds", "42、43、44"],
            ["參與端／攻擊者", "5 個；client 0 為攻擊者"],
            ["每次訓練", "10 輪，每端每輪訓練一遍自己的資料"],
            ["最佳化", "SGD；learning rate = 0.01；batch size = 32"],
            ["評估", "同一份 MNIST 測試集；主要比較第 10 輪"],
            [
                "配對控制",
                "同一 seed 下，Clean／Poisoned 的初始模型、資料分配及洗牌 seed 相同",
            ],
        ],
    )
    p(
        "A、B 各有五列設定，但 Alpha = 0.1、Scale = 10 是交會點。因此有九個不同設定，每個設定為三個 seeds × 兩種條件，共 54 次訓練。不是 60 次獨立訓練；B 中重複的 Clean 結果也不是額外的獨立樣本。"
    )
    sub("二、每輪流程")
    p(
        "全域模型送到五個 client → 各端用自己的資料訓練 → client 0 在攻擊條件下放大更新 → 依樣本數加權合併 → 使用測試集計算準確率與 Loss → 進入下一輪。"
    )
    sub("三、資料分配")
    p(
        "對每一類數字，用 Dirichlet 分布產生五個分配比例，再將該類圖片分到五個 client。Alpha 控制比例差異。每張訓練圖片只分到一個 client；同一 seed 的 Clean／Poisoned 使用完全相同的分配。這個方法不保證五端樣本數相等，所以 Alpha 同時影響資料種類與資料量。"
    )
    page("參、研究過程及方法（續）")
    sub("四、模型更新與合併公式")
    formula(
        "Δw = w_local − w_start",
        "w_start 是本輪開始的共同模型；w_local 是某個 client 訓練後的模型；Δw 是它這一輪學到的變化。",
    )
    formula(
        "w_send = w_start + s × Δw",
        "此式用於攻擊者，s 代表 Scale。Scale = 10 是把更新放大十倍，不是把整個模型直接乘十倍。",
    )
    formula(
        "w_next = Σ [(n_k / N) × w_send,k]",
        "Σ 表示把五個 client 的結果加起來；n_k 是第 k 個 client 的訓練樣本數，w_send,k 是該端送出的模型；N = 60,000。資料較多的 client，在 FedAvg 中的權重較大。",
    )
    sub("五、如何計算攻擊損害")
    formula(
        "Accuracy = 正確辨識張數 / 測試圖片總數",
        "程式以 0～1 記錄準確率；圖表乘以 100 後顯示百分比。",
    )
    formula(
        "Dacc = 100 × (Accuracyclean − Accuracypoisoned)",
        "Dacc 的單位是百分點。例如 90% 降到 70%，損害為 20 個百分點，不是相對下降 20%。",
    )
    formula(
        "Dloss = Losspoisoned − Lossclean",
        "兩種損害都以正值表示變差；負值表示該次結果略有改善，不應改成零。",
    )
    formula(
        "平均損害 = (D₄₂ + D₄₃ + D₄₄) / 3",
        "先算每個 seed 的配對差，再平均。誤差棒使用三個配對差的樣本標準差；不是信賴區間，也不是把十輪當成十次重複。",
    )
    formula(
        "樣本標準差 = √[ Σ (Dᵢ − 平均損害)² / (3 − 1) ]",
        "標準差越大，代表三個 seeds 的結果越分散。只有三次重複時，估計仍可能不穩定。",
    )
    for ex, title in [
        ("A", "肆、研究結果｜實驗 A：改變 Alpha"),
        ("B", "肆、研究結果｜實驗 B：改變 Scale"),
    ]:
        page(title)
        rows = [r for r in summary if r["experiment"] == ex]
        var = "alpha" if ex == "A" else "scale"
        table(
            [var.title(), "Clean（%）", "Poisoned（%）", "損害（百分點）"],
            [
                [
                    f"{float(r[var]):g}",
                    f"{float(r['clean_accuracy_mean'])*100:.2f} ± {float(r['clean_accuracy_sample_std'])*100:.2f}",
                    f"{float(r['poisoned_accuracy_mean'])*100:.2f} ± {float(r['poisoned_accuracy_sample_std'])*100:.2f}",
                    f"{float(r['accuracy_damage_pp_mean']):.2f} ± {float(r['accuracy_damage_pp_sample_std']):.2f}",
                ]
                for r in rows
            ],
        )
        p("表中為平均 ± 樣本標準差；每列使用三個 seeds。所有主要比較取第 10 輪。")
        picture(
            f"{ex}_overview",
            f'圖 {1 if ex=="A" else 2}　實驗 {ex} 的準確率、Loss 與配對損害。橫軸依設定等距排列；誤差棒可能跨過零，不代表原始 Loss 為負。',
        )
        if ex == "A":
            p(
                "無攻擊準確率隨 Alpha 增大而提高，且在 Alpha ≥ 1 時約為 90%～91%。平均準確率損害也隨 Alpha 增大，從 26.76 增至 64.68 個百分點。Loss 損害並未呈現相同的單調趨勢，因此不能說每個指標都隨 Alpha 增加。"
            )
        else:
            p(
                "Scale = 1 的準確率損害為零，平均 Loss 差約 1.07 × 10⁻⁸，符合浮點運算誤差的量級。Scale 從 5 到 10 時，平均準確率損害由 5.00 增至 34.22 個百分點；Scale = 20 時為 40.05。但這不足以證明所有情況都存在相同的臨界倍率。"
            )
    page("肆、研究結果｜各 seed 與訓練過程")
    picture(
        "seed_damage",
        "圖 3　三個 seeds 各自的準確率損害；平均趨勢不代表每個 seed 都有相同趨勢。",
    )
    p(
        "實驗 A 的平均損害逐步增加，但 seed 42 在 Alpha = 1 的損害比 Alpha = 0.1 小，seed 43 在 Alpha = 10 也比 Alpha = 1 小。實驗 B 的 seed 43 對放大較不敏感，Scale = 20 時損害仍只有 0.48 個百分點；seed 44 則從 Scale = 10 的 46.38 降到 Scale = 20 的 43.52。因此「Scale 越大」是本組平均值的趨勢，不是每次都成立的定律。"
    )
    picture(
        "training_curves",
        "圖 4　實驗 B 的平均訓練曲線。Clean 與 Scale = 1 幾乎重疊；曲線只顯示平均，不取代各 seed 的檢查。",
    )
    page("伍、討論｜攻擊者資料量的影響")
    picture(
        "attacker_weight",
        "圖 5　實驗 A 中 client 0 的資料占比與損害。每個點是一個 alpha／seed 組合，尚未控制其他因素。",
    )
    weightrows = [r for r in weights if r["experiment"] == "A" and r["alpha"] == 0.1]
    table(
        ["Seed", "攻擊者樣本數", "FedAvg 權重", "Scale = 10 損害"],
        [
            [
                r["seed"],
                r["attacker_samples"],
                f"{100*r['attacker_weight']:.2f}%",
                f"{r['accuracy_damage_pp']:.2f} 百分點",
            ]
            for r in weightrows
        ],
    )
    p(
        "在 Alpha = 0.1 時，seed 43 的攻擊者只有 1,135 張資料，占全部訓練資料約 1.89%；seed 42 與 44 分別約占 23.50% 和 29.85%。這可能是 seed 43 受到較小影響的原因之一，但尚不能據此證明因果。各端數字種類、更新方向與訓練狀態也不同。"
    )
    formula(
        "本輪攻擊者的更新係數 = s × n_attack / N",
        "n_attack 是攻擊者的樣本數。在單次 FedAvg 合併中，攻擊者的更新先乘 Scale，再乘樣本占比。這能說明兩者為何都值得記錄，但不能直接用這個係數預測十輪後的準確率。",
    )
    page("伍、討論｜可信度與限制")
    sub("一、已完成的檢查")
    p(
        "目前報告的產生腳本已核對九個設定的原始檔案雜湊、60,000 張訓練圖片的完整分配、Clean／Poisoned 配對識別資料，並從每次訓練的最後一輪重新計算 per_seed.csv 與 summary.csv。報告沒有重新訓練模型。"
    )
    p(
        "先前已修正兩項程式錯誤：第一，訓練前參數必須複製，避免被後續訓練改寫；第二，資料分配的最後邊界必須補到該類資料末端，避免浮點數捨去造成遺漏。舊錯誤版本的 poisoning 結果不納入。"
    )
    sub("二、目前不能下的結論")
    for t in [
        "只有三個 seeds，且多組標準差很大，尚未建立統計顯著性。",
        "只測 MNIST、一種小型神經網路、五個 client 與十輪訓練，不能推論所有資料或模型。",
        "Alpha 同時改變標籤差異與樣本數，尚未單獨測出資料種類差異的影響。",
        "兩組掃描不是所有 Alpha × Scale 的組合，尚未完整測試兩者互相影響的情況。",
        "目前只有更新放大攻擊，沒有實作防禦方法、後門或 Sybil 攻擊。",
        "此為本機依序模擬，不是實際分散式網路效能測試。",
    ]:
        p("• " + t)
    sub("三、為什麼要同時看準確率與 Loss？")
    p(
        "準確率只看答案是否正確，Loss 還會受到模型對答案有多確定的影響。例如模型很有把握地答錯，Loss 可能明顯上升。因此兩項指標不一定同時以相同比例變動，不能只選擇符合預期的指標。"
    )
    page("陸、結論與後續工作")
    sub("一、目前結論")
    for t in [
        "在本次十輪實驗中，較大的 Alpha 對應較高且較穩定的 Clean 準確率。",
        "固定 Scale = 10 時，平均準確率損害隨 Alpha 增大；本結果不支持「越 Non-IID 一定越容易受攻擊」的簡單說法。",
        "固定 Alpha = 0.1 時，Scale 越大，平均準確率損害與 Loss 損害越大；但各 seed 並非完全一致。",
        "攻擊者樣本占比值得進一步控制。現有資料能提供研究方向，還不能解釋所有原因。",
    ]:
        p("• " + t)
    sub("二、下一階段計畫（尚未執行）")
    p(
        "先增加不同 seeds，確認結果是否穩定。接著設計每個 client 樣本數相同、只改變數字種類分布的分配方式，檢查 Alpha 的影響是否仍然存在。再考慮交換攻擊者 client 編號、增加訓練輪數，以及測試限制更新大小等防禦方法。每次只改變一個主要條件，並保留配對的 Clean 結果。"
    )
    sub("三、研究進度")
    table(
        ["工作", "狀態"],
        [
            ["實驗 A、B 的三個 seeds", "已完成；九個不同設定、54 次訓練"],
            ["原始資料核對與圖表", "已完成，本腳本可重建"],
            ["控制攻擊者資料量", "待設計"],
            ["更多 seeds／資料集／模型", "待執行"],
            ["防禦、後門、Sybil", "尚未完成"],
        ],
    )
    page("柒、參考資料與附錄")
    p(
        "［1］McMahan, B., Moore, E., Ramage, D., Hampson, S., & Agüera y Arcas, B.（2017）. Communication-Efficient Learning of Deep Networks from Decentralized Data. PMLR 54, 1273–1282. https://proceedings.mlr.press/v54/mcmahan17a.html （用於聯邦學習與模型平均的背景說明。）"
    )
    p(
        "［2］PyTorch / torchvision. MNIST dataset documentation. https://docs.pytorch.org/vision/stable/generated/torchvision.datasets.MNIST.html （查閱日期：2026-09-21；用於資料載入介面說明。）"
    )
    p(
        "［3］本研究程式與資料：FL-project 的 client.py、run_matched.py、run_study.py，以及 results/studies/alpha_scale/。報告數值以保存的實驗結果為依據。"
    )
    p(
        "［4］使用者提供的 TISF.docx。僅參考報告形式與章節安排，未沿用其中駕駛偵測的實驗內容、作者資料或成果。此報告不表示已符合最新正式參賽格式。"
    )
    sub("附錄一：實驗 B 的三個 seeds 詳細損害")
    table(
        ["Scale", "Seed 42", "Seed 43", "Seed 44"],
        [
            [
                f"{s:g}",
                *[
                    f"{next(r['accuracy_damage_pp'] for r in computed if r['experiment']=='B' and r['scale']==s and r['seed']==seed):.2f}"
                    for seed in [42, 43, 44]
                ],
            ]
            for s in [1, 2, 5, 10, 20]
        ],
    )
    p("單位：百分點。負值保留原始方向，不改成零。")
    sub("附錄二：重建圖表與報告")
    p("在專案根目錄執行：\n.venv/bin/python scripts/build_science_report.py")
    p(
        "輸出資料夾：reports/current/。包含 Word 報告、PNG／PDF 圖表、攻擊者權重資料表及 validation.json。執行腳本只讀取實驗資料，不會啟動訓練，也不會修改 TISF.docx。"
    )
    # 第四階段：保存 Word，重新開啟確認圖表與文件結構完整。
    path = out / "聯邦學習_目前研究報告.docx"
    doc.save(path)
    # Verify the saved package and key content, rather than relying only on save success.
    check = Document(path)
    assert len(check.inline_shapes) == 6 and len(check.tables) >= 9
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
    (out / "README.md").write_text(
        "# 目前研究報告\n\n- 聯邦學習_目前研究報告.docx：可編輯 Word 報告。\n- figures/：6 張 PNG 圖及對應 PDF，適合簡報或列印。\n- attacker_weight_and_damage.csv：每個設定／seed 的攻擊者樣本占比及損害。\n- validation.json：資料核對紀錄與輸入檔案雜湊。\n\n重建：`.venv/bin/python scripts/build_science_report.py`。需 matplotlib、numpy、python-docx 與中文字型。\n未記錄的學校、作者、指導老師與參賽編號留待填寫。\n",
        encoding="utf-8",
    )
    print(
        f"Created {path}; {len(unique)} settings / {len(unique)*6} training runs verified; 6 figures (PNG + PDF)."
    )


if __name__ == "__main__":
    main()
